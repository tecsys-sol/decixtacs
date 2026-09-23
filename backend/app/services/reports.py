"""Module 18 - reports (daily/weekly/monthly) exported as CSV, Excel or PDF."""

from __future__ import annotations

import csv
import io
import uuid
from dataclasses import dataclass
from datetime import datetime, timedelta

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db.base import utcnow
from app.models import (
    AuditEvent,
    CommandLog,
    ConfigBackup,
    Device,
    DeviceComplianceScore,
    ComplianceRun,
    LoginHistory,
    TacacsAuthEvent,
)

PERIODS = {"daily": timedelta(days=1), "weekly": timedelta(days=7), "monthly": timedelta(days=30)}
REPORT_TYPES = ("device_changes", "config_changes", "user_activity", "compliance", "tacacs")


@dataclass
class Report:
    title: str
    columns: list[str]
    rows: list[list]
    start: datetime
    end: datetime


def build(db: Session, tenant_id: uuid.UUID, report_type: str, period: str = "daily",
          end: datetime | None = None) -> Report:
    end = end or utcnow()
    start = end - PERIODS[period]
    if report_type in ("device_changes", "config_changes"):
        q = (select(ConfigBackup.collected_at, Device.hostname, ConfigBackup.author, ConfigBackup.reason,
                    ConfigBackup.lines_added, ConfigBackup.lines_removed, ConfigBackup.risk_score, ConfigBackup.commit_sha)
             .join(Device, Device.id == ConfigBackup.device_id)
             .where(ConfigBackup.tenant_id == tenant_id, ConfigBackup.changed, ConfigBackup.collected_at.between(start, end))
             .order_by(ConfigBackup.collected_at))
        cols = ["Time", "Device", "Author", "Reason", "Added", "Removed", "Risk", "Commit"]
        rows = [list(r) for r in db.execute(q)]
        if report_type == "device_changes":
            agg: dict[str, list] = {}
            for r in rows:
                a = agg.setdefault(r[1], [r[1], 0, 0, 0, set()])
                a[1] += 1
                a[2] += r[4] or 0
                a[3] += r[5] or 0
                a[4].add(r[2] or "?")
            cols = ["Device", "Changes", "Lines added", "Lines removed", "Authors"]
            rows = [[a[0], a[1], a[2], a[3], ", ".join(sorted(a[4]))] for a in agg.values()]
    elif report_type == "user_activity":
        cmds = dict(db.execute(select(CommandLog.username, func.count()).where(
            CommandLog.tenant_id == tenant_id, CommandLog.timestamp.between(start, end)).group_by(CommandLog.username)).all())
        logins = dict(db.execute(select(LoginHistory.username, func.count()).where(
            LoginHistory.tenant_id == tenant_id, LoginHistory.success, LoginHistory.timestamp.between(start, end)
        ).group_by(LoginHistory.username)).all())
        fails = dict(db.execute(select(LoginHistory.username, func.count()).where(
            LoginHistory.tenant_id == tenant_id, LoginHistory.success.is_(False), LoginHistory.timestamp.between(start, end)
        ).group_by(LoginHistory.username)).all())
        audits = dict(db.execute(select(AuditEvent.actor_name, func.count()).where(
            AuditEvent.tenant_id == tenant_id, AuditEvent.timestamp.between(start, end)).group_by(AuditEvent.actor_name)).all())
        users = sorted(set(cmds) | set(logins) | set(fails) | set(audits))
        cols = ["User", "Portal logins", "Failed logins", "Device commands", "Portal actions"]
        rows = [[u, logins.get(u, 0), fails.get(u, 0), cmds.get(u, 0), audits.get(u, 0)] for u in users]
    elif report_type == "compliance":
        run = db.scalar(select(ComplianceRun).where(ComplianceRun.tenant_id == tenant_id, ComplianceRun.finished_at.is_not(None))
                        .order_by(ComplianceRun.started_at.desc()).limit(1))
        cols = ["Device", "Score", "Passed", "Failed"]
        rows = []
        if run:
            q = (select(Device.hostname, DeviceComplianceScore.score, DeviceComplianceScore.passed, DeviceComplianceScore.failed)
                 .join(Device, Device.id == DeviceComplianceScore.device_id)
                 .where(DeviceComplianceScore.run_id == run.id).order_by(DeviceComplianceScore.score))
            rows = [list(r) for r in db.execute(q)]
    elif report_type == "tacacs":
        q = (select(TacacsAuthEvent.kind, TacacsAuthEvent.result, func.count())
             .where(TacacsAuthEvent.tenant_id == tenant_id, TacacsAuthEvent.timestamp.between(start, end))
             .group_by(TacacsAuthEvent.kind, TacacsAuthEvent.result))
        cols = ["Type", "Result", "Count"]
        rows = [list(r) for r in db.execute(q)]
        acct = db.scalar(select(func.count()).select_from(CommandLog).where(
            CommandLog.tenant_id == tenant_id, CommandLog.timestamp.between(start, end)))
        rows.append(["accounting", "commands", acct or 0])
    else:
        raise ValueError(f"unknown report type {report_type}")
    title = f"{report_type.replace('_', ' ').title()} report ({period})"
    return Report(title, cols, rows, start, end)


def _cell(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    return "" if v is None else str(v)


def to_csv(r: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(r.columns)
    w.writerows([[_cell(c) for c in row] for row in r.rows])
    return buf.getvalue().encode()


def to_xlsx(r: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([r.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"{_cell(r.start)} - {_cell(r.end)} UTC"])
    ws.append([])
    ws.append(r.columns)
    for c in ws[4]:
        c.font = Font(bold=True)
    for row in r.rows:
        ws.append([_cell(c) if isinstance(c, datetime) else c for c in row])
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col[3:] if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(r: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), title=r.title)
    styles = getSampleStyleSheet()
    data = [r.columns] + [[_cell(c)[:60] for c in row] for row in r.rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
    ]))
    story = [Paragraph(r.title, styles["Title"]),
             Paragraph(f"{_cell(r.start)} - {_cell(r.end)} UTC - NetworkOps Manager", styles["Normal"]),
             Spacer(1, 12), table if r.rows else Paragraph("No data for this period.", styles["Normal"])]
    doc.build(story)
    return buf.getvalue()


RENDERERS = {"csv": (to_csv, "text/csv"),
             "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
             "pdf": (to_pdf, "application/pdf")}
